"""Read the unedited SIO workbooks and reconcile branch-year supply depths.

Workbook 6 reports irrigated area by branch. Workbook 8 reports supplied
water by site. Their Arabic labels are joined only after normalizing the
alif/hamza spelling variant; no geographic or many-to-one allocation is made.
"""
from __future__ import annotations

import argparse
import hashlib
import unicodedata
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "code/src"
OLD = ROOT / "code/src/sio_irrigation_2020_2022.csv"
NEW = ROOT / "code/src/sio_irrigation_2020_2022_v8.csv"
AUDIT = ROOT / "tables/sio_workbook_join_audit_v8.csv"
ALIASES = {
    "الاحساء": "Al-Ahsa", "الرياض": "Riyadh", "القطيف": "Qatif",
    "دومة الجندل": "Dumat Al-Jandal", "الافلاج": "Al-Aflaj",
    "المدينة المنورة": "Madinah",
}


def norm(value: str) -> str:
    text = unicodedata.normalize("NFKC", str(value)).strip()
    text = "".join(c for c in text if not unicodedata.combining(c))
    return " ".join(text.translate(str.maketrans("أإآٱ", "اااا")).split())


def number(value, *, no_flow_marker=False):
    if value == "*" and no_flow_marker:
        return 0.0
    if value is None or value == "":
        return np.nan
    if value == "*":
        raise ValueError("Unexpected * outside a source-water column")
    return float(value)


def sheet_rows(path: Path):
    sheet = load_workbook(path, read_only=True, data_only=True).active
    return list(sheet.iter_rows(values_only=True))


def parse_area(path: Path) -> pd.DataFrame:
    rows = sheet_rows(path)
    assert "السنة" in str(rows[1][5]) and "الفرع" in str(rows[1][6])
    out, year = [], None
    for excel_row, row in enumerate(rows[2:], 3):
        if row[5] is not None:
            year = int(row[5])
        if row[6] is None:
            continue
        branch = norm(row[6])
        out.append(dict(year=year, branch_key=branch, area_label=str(row[6]),
                        area_excel_row=excel_row, total_area_ha=number(row[7]),
                        irrigated_area_ha=number(row[8])))
    frame = pd.DataFrame(out)
    assert not frame.duplicated(["year", "branch_key"]).any()
    return frame


def parse_supply(path: Path) -> pd.DataFrame:
    rows = sheet_rows(path)
    assert "السنة" in str(rows[1][6]) and "الموقع" in str(rows[1][7])
    out, year = [], None
    for excel_row, row in enumerate(rows[2:], 3):
        if isinstance(row[6], (int, float)):
            year = int(row[6])
        if row[7] is None:
            continue
        branch = norm(row[7])
        vals = [number(row[i], no_flow_marker=True) for i in (8, 9, 10)]
        if any(np.isnan(v) for v in vals):
            raise ValueError(f"Blank source volume in 8.xlsx row {excel_row}; not a no-flow marker")
        out.append(dict(year=year, branch_key=branch, site_label=str(row[7]),
                        supply_excel_row=excel_row, reclaimed_m3=vals[0],
                        groundwater_m3=vals[1], agri_drainage_m3=vals[2]))
    frame = pd.DataFrame(out)
    assert not frame.duplicated(["year", "branch_key"]).any()
    return frame


def build(area: pd.DataFrame, supply: pd.DataFrame, old: pd.DataFrame):
    joined = area.merge(supply, on=["year", "branch_key"], how="outer",
                        validate="one_to_one", indicator=True).sort_values(["year", "branch_key"])
    if not set(joined.branch_key).issubset(ALIASES):
        raise ValueError("Unmapped Arabic label: " + str(set(joined.branch_key) - set(ALIASES)))
    joined["branch"] = joined.branch_key.map(ALIASES)
    joined["supplied_m3"] = joined[["reclaimed_m3", "groundwater_m3", "agri_drainage_m3"]].sum(axis=1,
                                                                                              min_count=3)
    can_depth = (joined._merge.eq("both") & joined.irrigated_area_ha.gt(0) & joined.supplied_m3.gt(0))
    joined["depth_m"] = np.where(can_depth, joined.supplied_m3 / joined.irrigated_area_ha / 10000, np.nan)
    old = old.copy()
    old["branch_key"] = old.branch_ar.map(norm)
    if old.duplicated(["year", "branch_key"]).any():
        raise ValueError("Duplicate saved branch-year")
    old["old_supplied_m3"] = old[["reclaimed_m3", "groundwater_m3", "agri_drainage_m3"]].fillna(0).sum(axis=1)
    old["old_depth_m"] = old.old_supplied_m3 / old.irrigated_area_ha / 10000
    joined = joined.merge(old[["year", "branch_key", "old_supplied_m3", "old_depth_m"]],
                          on=["year", "branch_key"], how="left", validate="one_to_one")
    joined["depth_difference_m"] = joined.depth_m - joined.old_depth_m
    joined["same_branch_year_units"] = can_depth
    joined["in_named_jeddah_makkah_branch"] = False
    valid = joined[can_depth].copy()
    if len(valid) != len(old) or valid.old_depth_m.isna().any():
        raise ValueError("New and saved branch-year populations differ; inspect audit before use")
    for field in ("reclaimed_m3", "groundwater_m3", "agri_drainage_m3", "irrigated_area_ha"):
        reference = old.set_index(["year", "branch_key"])[field].fillna(0)
        current = valid.set_index(["year", "branch_key"])[field].fillna(0)
        if not np.allclose(reference.reindex(current.index), current, rtol=0, atol=1e-9):
            raise ValueError(f"Workbook/source discrepancy in {field}")
    result = valid[["year", "branch", "area_label", "total_area_ha", "irrigated_area_ha",
                    "reclaimed_m3", "groundwater_m3", "agri_drainage_m3"]].rename(columns={"area_label": "branch_ar"})
    for field in ("reclaimed_m3", "groundwater_m3", "agri_drainage_m3"):
        result[field] = result[field].replace(0, np.nan)
    return result, joined


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw", type=Path, default=RAW)
    args = parser.parse_args()
    if NEW.exists() or AUDIT.exists():
        raise SystemExit("Refusing to overwrite v8 SIO outputs")
    six, eight = args.raw / "6.xlsx", args.raw / "8.xlsx"
    if not six.is_file() or not eight.is_file():
        raise SystemExit("Both unedited workbooks are required")
    hashes = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in (six, eight)}
    area, supply = parse_area(six), parse_supply(eight)
    result, audit = build(area, supply, pd.read_csv(OLD))
    result.to_csv(NEW, index=False)
    audit.to_csv(AUDIT, index=False)
    print("SIO originals", hashes)
    print("area rows", len(area), "site rows", len(supply), "valid branch-years", len(result))
    print("site-year rows without an area-year", int(audit._merge.eq("right_only").sum()))
    print("unmatched area rows", int(audit._merge.eq("left_only").sum()))
    print("max absolute saved-depth difference m", float(audit.depth_difference_m.abs().max()))
    water = pd.read_csv(ROOT / "tables/water.csv")
    lo, hi = float(water.depth_m.min()), float(water.depth_m.max())
    print("outside assumed scenario band", int((~audit.loc[audit.depth_m.notna(), "depth_m"].between(lo, hi)).sum()),
          "of", int(audit.depth_m.notna().sum()), "band", lo, hi)


if __name__ == "__main__":
    main()
